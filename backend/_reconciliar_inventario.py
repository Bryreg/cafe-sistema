# -*- coding: utf-8 -*-
"""Reconcilia inventario (Excel lleno) + menú (precios) + conteo diario → preview + JSON.
NO toca la base. Genera INVENTARIO_VIDA_carga.csv (revisión) e inventario_inicial.json."""
import openpyxl, unicodedata, re, json, csv
from collections import defaultdict, Counter

def norm(s): return re.sub(r'\s+', ' ', unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()).strip().upper()
def key(s): return re.sub(r'[^A-Z0-9]', '', norm(s))

XLSX = 'C:/Users/bmgpe/Downloads/6. FORMATO INVENTARIO CIERRE DE MES - VIDA.xlsx'
wb = openpyxl.load_workbook(XLSX, data_only=True); ws = wb[wb.sheetnames[0]]

MERGE = {
 'CAFEALTATOSTIONX2500GR': 'Cafe Alta Tostion (granel)', 'CAFEALTATOSTION': 'Cafe Alta Tostion (granel)',
 'AZUCARAGRANELX2500GR': 'Azucar a Granel', 'AZUCARAGRANELX2500': 'Azucar a Granel', 'AZUCARAGRANELX2501': 'Azucar a Granel',
 'SABORIZANTECANELA': 'Saborizante Canela', 'SABORIZANTECINNAMONCANELA': 'Saborizante Canela',
 'SABORIZANTEFRESAKIWI': 'Saborizante Kiwi Fresa', 'SABORIZANTEKIWIFRESA': 'Saborizante Kiwi Fresa',
 'MATCHA': 'Matcha', 'TEMATCHA': 'Matcha',
 'SALSAFRUTOSROJOS': 'Salsa Frutos Rojos', 'SALSASUNDAEFRUTOSROJOS': 'Salsa Frutos Rojos',
 'SALSASUNDAEDECARAMELO': 'Salsa Caramelo', 'SALSASUNDAEDECHOCOLATE': 'Salsa Chocolate',
 'VASO12OZ': 'Vaso Carton 12oz', 'VASOCARTON12OZ': 'Vaso Carton 12oz',
 'VASO16OZ': 'Vaso Carton 16oz', 'VASOCARTON16OZ': 'Vaso Carton 16oz',
 'HELADOGOGOSCHOCOLATE': 'Helado Chocolate', 'HELADOGOGOSVAINILLA': 'Helado Vainilla',
 'LICORAMARETTOX750ML': 'Licor Amaretto x750ml', 'LICORAMARETTO': 'Licor Amaretto x750ml',
 'LICORWHISKEYX700ML': 'Licor Whisky Black & White x700ml', 'LICORWHISKYBLACKWHITEX700ML': 'Licor Whisky Black & White x700ml',
 'LICORWHISKEYX1000ML': 'Licor Whisky Black & White x1000ml', 'LICORWHISKYBLACKWHITEX1000ML': 'Licor Whisky Black & White x1000ml',
 'LICORBAILEYSX700ML': 'Licor Baileys x700ml', 'LICORBAILEYSX1000ML': 'Licor Baileys x1000ml',
}
def canon(d): return MERGE.get(key(d), str(d).strip())

SELL = [(2, 1, 3), (6, 5, 7), (10, 9, 11)]; ABI = (15, 14, 16)
sell = defaultdict(float); abi = defaultdict(float); names = {}
def add(store, d, q):
    n = canon(d); k = key(n); store[k] += q; names[k] = n
for row in ws.iter_rows(values_only=True):
    for d, m, q in SELL:
        de = row[d] if d < len(row) else None; qq = row[q] if q < len(row) else None
        if de and str(de).strip() and isinstance(qq, (int, float)) and norm(de) != 'DESCRIPCION':
            add(sell, str(de), float(qq))
    d, m, q = ABI
    de = row[d] if d < len(row) else None; qq = row[q] if q < len(row) else None
    if de and str(de).strip() and isinstance(qq, (int, float)) and norm(de) != 'DESCRIPCION':
        add(abi, str(de), float(qq))
prods = sorted(set(sell) | set(abi))

# Precios del menú: parsear PRODUCTOS_POR_GRUPOS directo (3 bloques de columnas).
precio_by = {}
try:
    mwb = openpyxl.load_workbook('C:/Users/bmgpe/Downloads/PRODUCTOS_POR_GRUPOS (2).xlsx', data_only=True)
    mws = mwb['PRODUCTOS']
    for row in mws.iter_rows(values_only=True):
        for (pc, prc) in [(1, 2), (5, 6), (9, 10)]:
            prod = row[pc] if pc < len(row) else None
            pr = row[prc] if prc < len(row) else None
            if prod and str(prod).strip() and isinstance(pr, (int, float)) and norm(prod) != 'PRODUCTO':
                precio_by.setdefault(key(prod), int(pr))
except Exception as e:
    print('WARN precios:', e)

PAST = ['TORTA', 'CROISSANT', 'CROISANT', 'OMELET', 'PASTEL', 'ALMOJABANA', 'PALITO DE QUESO', 'MASA DE PANDEBONO']
INS = ['VASO', 'TAPA', 'BOLSA', 'SERVILLETA', 'PITILLO', 'MEZCLADOR', 'GUANTE', 'JABON', 'DETERGENTE', 'LIMPIA', 'BLANQUEADOR',
       'CITRONELA', 'ALCOHOL', 'PANO', 'WYPALL', 'ESPONJILLA', 'MALLA', 'ROLLO', 'PAPEL', 'GEL', 'CAJA', 'ENDULZANTE',
       'TENEDOR', 'CUCHILLO', 'JARRA', 'TAZA', 'COPA', 'PLATO', 'ESCOBA', 'RECOGEDOR', 'TRAPEADOR', 'FILTRO', 'CUCHARA', 'AZUCAR BLANCA TUBOS']
BOTELLA = ['SABORIZANTE', 'LICOR', 'SALSA', 'SOUR CREAM']
BOLSA_FR = ['CAFE ALTA', 'AZUCAR A GRANEL', 'CANELA', 'COCOA', 'MILO', 'CHAI', 'MATCHA', 'LECHE EN POLVO',
            'LECHE CONDENSADA', 'GALLETA OREO', 'PANELA', 'HELADO', 'MEZCLA GRANIZADO', 'ESPRESSOS FRIOS']
NO_FRAC = ['LECHE ENTERA', 'AROMATICA', 'PULPA', 'CAFE LIBRA', 'BATI CREMA', 'AZUCAR BLANCA TUBOS']

def cat(n):
    x = norm(n)
    if any(k in x for k in PAST): return 'pasteleria'
    if any(k in x for k in INS): return 'insumo'
    return 'bebida'
def frac_env(n):
    x = norm(n)
    if any(k in x for k in NO_FRAC): return (False, None)
    if any(k in x for k in BOTELLA): return (True, 'botella')
    if any(k in x for k in BOLSA_FR): return (True, 'bolsa')
    return (False, None)

DAILY = ['CAFEALTATOSTIONGRANEL', 'CAFELIBRAMEDIUMCAFEEXPORTACION', 'ALMOJABANAS', 'TORTACHOCOLATE', 'TORTANARANJA',
 'TORTAZANAHORIA', 'TORTAREDVELVET', 'PASTELDEPOLLO', 'PALITODEQUESO', 'OMELETTE', 'OMELETTEDEJAMONYQUESO',
 'CROISSANTQUESO', 'CROISSANTDECHOCOLATE', 'MASADEPANDEBONOX25GRAMOS', 'CHAI', 'AGUAMEDIUMBOTELLA',
 'AGUAMEDIUMCONGASBOTELLA', 'LICORBAILEYSX700ML', 'LICORBAILEYSX1000ML', 'LICORWHISKYBLACKWHITEX700ML',
 'LICORWHISKYBLACKWHITEX1000ML', 'LICORWHISKEYX700ML', 'LICORWHISKEYX1000ML', 'LICORAMARETTOX750ML',
 'SABORIZANTEVAINILLA', 'SABORIZANTECANELA', 'SABORIZANTEMACADAMIA', 'SABORIZANTEFRUTOSAMARILLOS', 'SABORIZANTEKIWIFRESA',
 'PULPADELULO', 'PULPADEMANGO', 'PULPADEMORA', 'PULPADELIMON', 'LECHEENTERAYDESLACTOSADA', 'LECHECONDENSADA',
 'SOURCREAM', 'LECHEENPOLVO', 'SALSACHOCOLATE', 'SALSACARAMELO', 'SALSAFRUTOSROJOS', 'SALSAMARACUYA',
 'CREMACHANTILLY', 'AZUCARBLANCATUBOS', 'AZUCARAGRANEL', 'MILO', 'GALLETAOREO', 'BATICREMA']
dset = set(DAILY)
PACK = {'CAFEALTATOSTIONGRANEL': 2500, 'LECHECONDENSADA': 300}

rows = []
for k in prods:
    n = names[k]; s = sell.get(k, 0); a = abi.get(k, 0)
    c = cat(n)
    # Fraccionable SOLO para ingredientes a granel (categoría bebida). Equipos/desechables (insumo) nunca.
    fr, env = frac_env(n) if c == 'bebida' else (False, None)
    diario = key(n) in dset
    precio = precio_by.get(key(n), 0)
    # Colisión ingrediente/bebida: Cocoa y Chai Latte existen como insumo Y como bebida de venta.
    if key(n) in ('COCOA', 'CHAILATTE'):
        precio = 0  # lo dejo como ingrediente; la bebida de venta la maneja el menú aparte
    # Preparados en uso (sin selladas): el valor pesado era lo abierto → 0 selladas, nivel por dibujo.
    EN_USO = {'ESPRESSOSFRIOS', 'MEZCLAGRANIZADO', 'LECHEENPOLVO'}
    if key(n) in EN_USO:
        sealed = 0
    elif key(n) in PACK:
        sealed = round(s / PACK[key(n)])
    else:
        sealed = round(s, 2)
    rows.append({'nombre': n, 'categoria': c, 'fraccionable': fr, 'envase': env,
                 'diario': diario, 'precio': precio, 'sellado': sealed, 'abierto_g': a})

with open('C:/Users/bmgpe/Desktop/INVENTARIO_VIDA_carga.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(['Producto', 'Categoria', 'Fraccionable', 'Envase', 'ConteoDiario', 'PrecioVenta', 'StockSellado'])
    for r in rows:
        w.writerow([r['nombre'], r['categoria'], 'si' if r['fraccionable'] else '', r['envase'] or '',
                    'si' if r['diario'] else '', r['precio'] or '', r['sellado']])
json.dump(rows, open('inventario_inicial.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

print('TOTAL productos:', len(rows))
print('Categoria:', dict(Counter(r['categoria'] for r in rows)))
print('Fraccionables:', sum(1 for r in rows if r['fraccionable']),
      '| Conteo diario:', sum(1 for r in rows if r['diario']),
      '| Con precio venta:', sum(1 for r in rows if r['precio']))
print('CSV -> Desktop/INVENTARIO_VIDA_carga.csv | JSON -> backend/inventario_inicial.json')
print('\n--- FRACCIONABLES (dibujo):')
for r in rows:
    if r['fraccionable']:
        print('  %-38s [%s] sellado=%g%s' % (r['nombre'], r['envase'], r['sellado'], '  DIARIO' if r['diario'] else ''))
print('\n--- CON PRECIO DE VENTA:')
for r in rows:
    if r['precio']:
        print('  %-38s $%s  sellado=%g' % (r['nombre'], format(r['precio'], ','), r['sellado']))
